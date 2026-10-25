import argparse
import json
from datetime import date, datetime
from pathlib import Path

from app.application.importation.use_cases import import_catalog
from app.adapters.outbound.persistence.importer import SqlCatalogImporter
from app.infrastructure.config import get_settings
from app.infrastructure.database import create_database


def read_catalog(path):
    if path.suffix.lower() == ".json":
        return json.loads(path.read_text())
    from openpyxl import load_workbook
    sheet = load_workbook(path, data_only=True)["Carreras"]
    headers = [c.value for c in sheet[1]]
    return [{k: v.isoformat() if isinstance(v, (date, datetime)) else v
             for k, v in zip(headers, row)}
            for row in sheet.iter_rows(min_row=2, values_only=True) if any(row)]


def main():
    parser = argparse.ArgumentParser(description="Carga idempotente del catálogo ORIENTA")
    parser.add_argument("path", type=Path)
    parser.add_argument("--editorial", type=Path, default=Path("data/contenido_editorial.json"))
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()
    engine, factory = create_database(get_settings().database_url)
    try:
        with factory.begin() as session:
            report = import_catalog(SqlCatalogImporter(session), read_catalog(args.path),
                                    json.loads(args.editorial.read_text()), args.dry_run)
        print(json.dumps(report, ensure_ascii=False))
    finally:
        engine.dispose()


if __name__ == "__main__":
    main()
